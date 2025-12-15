"""
XLSX Writer - Generate XLSX output files from extracted IDML data

Creates XLSX files with standard structure:
- Sheet "prodotti": Product-level metadata
- Sheet "sku": SKU-level technical specs
"""

from pathlib import Path
from typing import Dict, List, Optional, Any
from dataclasses import dataclass, field
import logging
from datetime import datetime

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None

from .xlsx_loader import PRODOTTI_COLUMNS, SKU_COLUMNS
from .table_extractor import SKUSpecs
from .text_extractor import ProductInfo

logger = logging.getLogger(__name__)


@dataclass
class ExtractionResult:
    """Complete extraction result from an IDML file"""
    source_file: str
    product_info: ProductInfo
    sku_specs: List[SKUSpecs]
    extraction_time: str = ""

    def __post_init__(self):
        if not self.extraction_time:
            self.extraction_time = datetime.now().isoformat()


class XLSXWriter:
    """Write extraction results to XLSX files"""

    # Styling
    HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    HEADER_FONT = Font(bold=True, color="FFFFFF")
    BORDER = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    def __init__(self, output_path: Path):
        if openpyxl is None:
            raise ImportError("openpyxl is required. Install with: pip install openpyxl")
        self.output_path = Path(output_path)
        self.workbook = Workbook()

    def write(self, results: List[ExtractionResult]) -> Path:
        """Write extraction results to XLSX file"""
        # Remove default sheet
        if 'Sheet' in self.workbook.sheetnames:
            del self.workbook['Sheet']

        # Create prodotti sheet
        self._write_prodotti_sheet(results)

        # Create sku sheet
        self._write_sku_sheet(results)

        # Save workbook
        self.workbook.save(self.output_path)
        logger.info(f"Wrote XLSX to: {self.output_path}")

        return self.output_path

    def _write_prodotti_sheet(self, results: List[ExtractionResult]):
        """Write product-level data"""
        ws = self.workbook.create_sheet("prodotti")

        # Write headers
        for col, header in enumerate(PRODOTTI_COLUMNS, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = self.HEADER_FILL
            cell.font = self.HEADER_FONT
            cell.border = self.BORDER
            cell.alignment = Alignment(horizontal='center')

        # Write data rows
        for row_idx, result in enumerate(results, 2):
            info = result.product_info
            info_dict = info.to_dict()

            for col, header in enumerate(PRODOTTI_COLUMNS, 1):
                value = info_dict.get(header, "")
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.border = self.BORDER

        # Auto-adjust column widths
        self._adjust_column_widths(ws)

    def _write_sku_sheet(self, results: List[ExtractionResult]):
        """Write SKU-level data"""
        ws = self.workbook.create_sheet("sku")

        # Write headers
        for col, header in enumerate(SKU_COLUMNS, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = self.HEADER_FILL
            cell.font = self.HEADER_FONT
            cell.border = self.BORDER
            cell.alignment = Alignment(horizontal='center')

        # Write data rows
        row_idx = 2
        for result in results:
            product_id = result.product_info.name

            for sku_specs in result.sku_specs:
                specs_dict = sku_specs.to_dict()
                specs_dict['product_id'] = product_id

                for col, header in enumerate(SKU_COLUMNS, 1):
                    # Map header to spec value
                    value = specs_dict.get(header, "")
                    cell = ws.cell(row=row_idx, column=col, value=value)
                    cell.border = self.BORDER

                row_idx += 1

        # Auto-adjust column widths
        self._adjust_column_widths(ws)

    def _adjust_column_widths(self, ws):
        """Auto-adjust column widths based on content"""
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)

            for cell in column:
                try:
                    cell_length = len(str(cell.value))
                    if cell_length > max_length:
                        max_length = cell_length
                except:
                    pass

            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width


def write_xlsx(results: List[ExtractionResult], output_path: Path) -> Path:
    """Convenience function to write XLSX file"""
    writer = XLSXWriter(output_path)
    return writer.write(results)


def create_extraction_result(
    source_file: str,
    product_info: ProductInfo,
    sku_specs: List[SKUSpecs]
) -> ExtractionResult:
    """Create an extraction result object"""
    return ExtractionResult(
        source_file=source_file,
        product_info=product_info,
        sku_specs=sku_specs
    )


def extract_and_write(idml_path: Path, output_path: Path) -> Path:
    """Extract from IDML and write to XLSX in one step"""
    from .idml_parser import parse_idml
    from .table_extractor import extract_specs_from_document
    from .text_extractor import extract_product_info

    # Parse IDML
    document = parse_idml(idml_path)

    # Extract data
    product_info = extract_product_info(document)
    sku_specs = extract_specs_from_document(document)

    # Create result
    result = create_extraction_result(
        source_file=str(idml_path.name),
        product_info=product_info,
        sku_specs=sku_specs
    )

    # Write XLSX
    return write_xlsx([result], output_path)
